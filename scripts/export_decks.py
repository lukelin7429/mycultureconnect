#!/usr/bin/env python3
"""Export a deck to PDF and PPTX, so a classroom PC can run it without a browser.

    python3 scripts/export_decks.py                 # every 115 interactive deck
    python3 scripts/export_decks.py beidou-jh-sdg2-zerohunger

Writes slides/<slug>/<slug>.pdf and slides/<slug>/<slug>.pptx.

Two things worth knowing:

* Click-to-reveal survives the export. A slide with a hidden answer is written
  out TWICE — once with the answer still hidden, once revealed — so advancing
  the PDF/PPTX reveals the answer exactly like tapping does on the web deck.
  The answer is never printed next to the question.
* The PPTX holds one full-bleed image per slide rather than editable text. That
  is deliberate: a school PC has no PingFang TC, and real text boxes would
  reflow and break the layout. Images look identical everywhere.

Needs headless-chrome (~/.local/bin, headless print), PyMuPDF and python-pptx.
"""
import os, re, sys, glob, time, subprocess, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Not the Chrome in /Applications: running the everyday browser headless leaves a
# bouncing icon in the Dock and sometimes a process that never exits.
CHROME = os.path.expanduser('~/.local/bin/headless-chrome')
CANVAS_W, CANVAS_H = 1200, 675
CHUNK = 4              # printed pages per Chrome run. Was 8 until the slides began enlarging
                       # their own text: bigger emoji make each page heavier, and a batch of
                       # eight started failing with no output. Smaller batches, same result.


def slide_blocks(html):
    """Pull out each <div class="deck-slide">…</div> with brace-free depth counting."""
    out, i = [], 0
    while True:
        m = re.compile(r'<div class="deck-slide"[^>]*>').search(html, i)
        if not m:
            return out
        depth, j = 1, m.end()
        while depth:
            nxt = re.compile(r'<div\b|</div>').search(html, j)
            if not nxt:
                return out
            depth += 1 if nxt.group(0) != '</div>' else -1
            j = nxt.end()
        out.append(html[m.end():j - len('</div>')])
        i = j


def reveal(block):
    """Same slide with every hidden answer shown."""
    b = re.sub(r'(<div class="[^"]*?)"([^>]*data-reveal-group)', r'\1 is-revealed"\2', block)
    b = re.sub(r'(class=")((?:[^"]*\s)?hs-tf)(["\s])', r'\1\2 is-revealed\3', b)
    b = re.sub(r'(class=")((?:[^"]*\s)?hs-hintbtn)(["\s])', r'\1\2 is-revealed\3', b)
    b = re.sub(r'(class=")((?:[^"]*\s)?hs-sortitem)(["\s])', r'\1\2 is-revealed\3', b)
    b = re.sub(r'(<[^>]*\bdata-reveal-group\b[^>]*class=")', r'\1is-revealed ', b)
    return b


def has_hidden_answer(block):
    return ('data-reveal-group' in block or 'hs-tf ' in block
            or 'hs-tf"' in block or 'hs-hintbtn' in block or 'hs-sortitem' in block)


PRINT_PAGE = """<!doctype html><meta charset="utf-8">
<link rel="stylesheet" href="file://{css}">
<style>
  @page {{ size: {w}px {h}px; margin: 0; }}
  html, body {{ margin:0; padding:0; background:#fff; }}
  /* the web page gets its font from body.deck-body; this body has no class,
     so without this line Chrome falls back to Times in every PDF and PPTX */
  body {{ font-family:var(--sans); }}
  .pg {{ width:{w}px; height:{h}px; overflow:hidden; page-break-after:always;
         display:flex; align-items:center; justify-content:center; background:#000; }}
  .pg:last-child {{ page-break-after:auto; }}
  /* the canvas is scaled by deck.js on the web; in print it is 1:1 */
  .hs-canvas {{ --hs-scale:1; transform:none !important; }}
  .pg > img {{ width:100%; height:100%; object-fit:contain; }}
  /* offline copies cannot play the video inline; the whole page links to it */
  .hs-ytcue {{ position:absolute; left:50%; bottom:16px; transform:translateX(-50%);
    background:rgba(0,0,0,.74); color:#fff; border-radius:999px; padding:8px 22px;
    font-size:21px; font-weight:700; white-space:nowrap; }}
  .hs-ytcue em {{ font-style:normal; font-family:var(--zh); font-weight:600; margin-left:8px; }}
  /* no .deck.anim-ready here, so every entrance animation is already finished */
</style>
{pages}
<!-- the same fit the live deck runs: each slide enlarges its text to fill the frame -->
<script src="file://{fit}"></script>
<script>hsFit(document);</script>
"""


def build_pages(slug):
    """One HTML page per printed page — a slide with a hidden answer yields two."""
    src = os.path.join(ROOT, 'slides', slug, 'index.html')
    html = open(src, encoding='utf-8').read()
    blocks = slide_blocks(html)
    if not blocks:
        raise SystemExit(f'{slug}: no slides found')
    pages, links, revealed = [], [], 0
    for b in blocks:
        yt = re.search(r'data-yt="([\w-]{11})"', b)
        if yt:
            # A PDF or PPTX page is only a picture of the video, so it must say
            # so and take you to it — a slide that looks playable but isn't
            # is how a class ends up waiting while the teacher hunts for the link.
            b = re.sub(r'(<button[^>]*class="hs-play"[^>]*>.*?</button>)',
                       r'\1<span class="hs-ytcue">▶ Click to play on YouTube'
                       r'<em>點一下到 YouTube 播放</em></span>', b, count=1)
        pages.append(b)
        links.append(f'https://www.youtube.com/watch?v={yt.group(1)}' if yt else None)
        if has_hidden_answer(b):
            pages.append(reveal(b))
            links.append(None)
            revealed += 1
    return pages, links, len(blocks), revealed


def write_print_html(dest, pages, tag):
    body = '\n'.join('<div class="pg">' + p + '</div>' for p in pages)
    # Drop the <picture> WebP sources so Chrome prints from the PNG. WebP is the
    # right choice over the network, but Chrome re-encodes it badly on the way
    # into a PDF — the same deck comes out ~40% larger.
    body = re.sub(r'<source[^>]*>', '', body)
    # /assets/... only resolves over http; make it absolute for file:// printing
    for attr in ('src', 'srcset', 'href'):
        body = body.replace(f'{attr}="/assets/', f'{attr}="file://{ROOT}/assets/')
    out = os.path.join(dest, f'print-{tag}.html')
    open(out, 'w', encoding='utf-8').write(PRINT_PAGE.format(
        css=os.path.join(ROOT, 'assets', 'css', 'deck.css'),
        fit=os.path.join(ROOT, 'assets', 'js', 'deck-fit.js'),
        w=CANVAS_W, h=CANVAS_H, pages=body))
    return out


def to_pdf(print_html, pdf_path, timeout=180):
    """Headless Chrome writes the PDF and then often refuses to exit on macOS,
    so wait for the file to stop growing and stop the process ourselves."""
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
    with tempfile.TemporaryDirectory() as prof:
        proc = subprocess.Popen([
            CHROME, '--headless=new', '--disable-gpu', '--no-first-run',
            '--user-data-dir=' + prof, '--allow-file-access-from-files',
            '--run-all-compositor-stages-before-draw',
            '--virtual-time-budget=20000', '--no-pdf-header-footer',
            '--print-to-pdf=' + pdf_path, 'file://' + print_html,
        ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        last, steady, deadline = -1, 0, time.time() + timeout
        while time.time() < deadline:
            if proc.poll() is not None:
                break
            size = os.path.getsize(pdf_path) if os.path.exists(pdf_path) else -1
            steady = steady + 1 if size > 0 and size == last else 0
            if steady >= 3:                     # 1.5 s with no growth = done
                break
            last = size
            time.sleep(0.5)
        if proc.poll() is None:
            proc.terminate()
            try:
                proc.wait(10)
            except subprocess.TimeoutExpired:
                proc.kill()
    if not os.path.exists(pdf_path) or os.path.getsize(pdf_path) < 1024:
        raise SystemExit('Chrome produced no PDF for ' + pdf_path)


# PowerPoint cannot play a YouTube link inside the slide: a click on the picture
# only opened the browser, and only in Slide Show. Where Dom's own finished
# animation is on this computer, the PPTX now carries the video itself, so it
# plays right on the slide with no internet. The copy is shrunk to 720p first
# (the 1080p masters are 50-90 MB each; this keeps a deck under ~15 MB more).
# Videos that are not here (general films, or a master not found) keep the link.
# the finished animations have lived in both places; the first that has the file wins
VIDEO_DIRS = [os.path.expanduser(p) for p in ("~/Downloads/Dom's animations/成品",
                                               "~/Downloads/工作檔案/4_學校教材/Dom's animations/成品")]
VIDEO_CACHE = os.path.expanduser('~/Library/Caches/mcc-deck-videos')
FFMPEG = next((p for p in ('/opt/homebrew/opt/ffmpeg-full/bin/ffmpeg', '/opt/homebrew/bin/ffmpeg', '/usr/local/bin/ffmpeg')
               if os.path.exists(p)), 'ffmpeg')
VIDEOS = {  # YouTube id -> finished file in one of VIDEO_DIRS
    'AbowvF_5KYk': 'SDG1_NoPoverty_bilingual_v1.mp4',
    '2SgHi7d4JUQ': 'SDG5_GenderEquality_bilingual_v1.mp4',
    'eZIfSZBYJ4I': 'Character_Honesty_bilingual_v1.mp4',
    '9Lkvbd8x_GI': 'Character_Kindness_bilingual_v1.mp4',
    'qQQhU5rkwXE': 'Character_Gratitude_bilingual_v1.mp4',
    'wr_tAh681ds': 'Character_AntiBullying_bilingual_v2.mp4',
    'z8pbdkLaWHI': 'SDG16_PeaceJustice_bilingual_v1.mp4',
    '13e7IjaoYJA': 'SDG11_SustainableCities_bilingual_v1.mp4',
    'KZ03usiTxZ4': 'SDG12_ResponsibleConsumption_bilingual_v1.mp4',
    '-eo8LoJJ7U4': 'SDG13_ClimateAction_bilingual_v1.mp4',
    'A9liiJZ3MVo': 'SDG6_CleanWater_bilingual_v1.mp4',
    'BCxuSzg532g': 'SDG4_QualityEducation_bilingual_v2.mp4',
    '4k-BTYRiWmo': 'SDG10_ReducedInequalities_bilingual_v1.mp4',
}


def local_video(url):
    """(small mp4, poster png) for a YouTube link, or None if we have no file for it."""
    m = re.search(r'v=([\w-]{11})', url or '')
    name = VIDEOS.get(m.group(1)) if m else None
    src = next((os.path.join(d, name) for d in VIDEO_DIRS if name and os.path.exists(os.path.join(d, name))), '')
    if not src:
        return None
    os.makedirs(VIDEO_CACHE, exist_ok=True)
    stem = os.path.join(VIDEO_CACHE, f'{m.group(1)}-{int(os.path.getmtime(src))}')
    mp4, poster = stem + '.mp4', stem + '.jpg'
    if not os.path.exists(mp4):
        subprocess.run([FFMPEG, '-v', 'error', '-y', '-i', src, '-vf', 'scale=1280:-2', '-c:v', 'libx264',
                        '-preset', 'slow', '-crf', '27', '-profile:v', 'high', '-pix_fmt', 'yuv420p',
                        '-c:a', 'aac', '-b:a', '96k', '-movflags', '+faststart', mp4], check=True)
    if not os.path.exists(poster):
        subprocess.run([FFMPEG, '-v', 'error', '-y', '-ss', '1.5', '-i', mp4, '-frames:v', '1', '-q:v', '3', poster], check=True)
    return mp4, poster


def to_pptx(pdf_path, pptx_path, links=(), scale=2.0):
    import fitz
    from pptx import Presentation
    from pptx.util import Inches, Emu

    doc = fitz.open(pdf_path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    embedded = 0
    with tempfile.TemporaryDirectory() as tmp:
        for n, page in enumerate(doc):
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale))
            # Text slides stay PNG (crisp, and small because the colors are flat);
            # a slide that is mostly a photograph is several times smaller as JPEG.
            png = os.path.join(tmp, f'{n:03d}.png')
            jpg = os.path.join(tmp, f'{n:03d}.jpg')
            pix.save(png)
            pix.save(jpg, jpg_quality=88)
            pick = jpg if os.path.getsize(jpg) < os.path.getsize(png) * 0.7 else png
            s = prs.slides.add_slide(blank)
            pic = s.shapes.add_picture(pick, 0, 0, prs.slide_width, prs.slide_height)
            url = links[n] if n < len(links) else None
            if not url:
                continue
            video = local_video(url)
            imgs = sorted(page.get_image_info(), key=lambda i: (i['bbox'][2] - i['bbox'][0]) * (i['bbox'][3] - i['bbox'][1]))
            if video and imgs:
                # the video sits exactly where the YouTube picture is on the page
                x0, y0, x1, y1 = imgs[-1]['bbox']
                kx, ky = prs.slide_width / page.rect.width, prs.slide_height / page.rect.height
                s.shapes.add_movie(video[0], Emu(int(x0 * kx)), Emu(int(y0 * ky)),
                                   Emu(int((x1 - x0) * kx)), Emu(int((y1 - y0) * ky)),
                                   poster_frame_image=video[1], mime_type='video/mp4')
                embedded += 1
            else:
                # no file for this one: in Slide Show a click on the picture opens it in the browser
                pic.click_action.hyperlink.address = url
    prs.save(pptx_path)
    if embedded:
        print(f'    {embedded} video(s) embedded in the PPTX')
    return len(doc)


def is_stale(slug):
    d = os.path.join(ROOT, 'slides', slug)
    src = os.path.join(d, 'index.html')
    outs = [os.path.join(d, slug + '.pdf'), os.path.join(d, slug + '.pptx')]
    if not all(os.path.exists(o) for o in outs):
        return True
    return any(os.path.getmtime(o) < os.path.getmtime(src) for o in outs)


def merge(parts, out, links=()):
    import fitz
    doc = fitz.open()
    for p in parts:
        with fitz.open(p) as part:
            doc.insert_pdf(part)
    for n, url in enumerate(links):
        if url:
            doc[n].insert_link({'kind': fitz.LINK_URI, 'from': doc[n].rect, 'uri': url})
    doc.save(out)
    doc.close()


def export(slug, chunk=CHUNK):
    out_dir = os.path.join(ROOT, 'slides', slug)
    pdf = os.path.join(out_dir, slug + '.pdf')
    pptx = os.path.join(out_dir, slug + '.pptx')
    with tempfile.TemporaryDirectory() as tmp:
        pages, links, n_slides, n_revealed = build_pages(slug)
        # Chrome's print pipeline gives up on a long, image-heavy document
        # ("Printing failed" with no output), so print in batches and stitch.
        parts = []
        for i in range(0, len(pages), chunk):
            part = os.path.join(tmp, f'part-{i:03d}.pdf')
            to_pdf(write_print_html(tmp, pages[i:i + chunk], i), part)
            parts.append(part)
        merge(parts, pdf, links)
        n_pages = to_pptx(pdf, pptx, links)
    print(f'  ✓ {slug}: {n_slides} slides (+{n_revealed} answer pages) '
          f'-> {n_pages}p  pdf {os.path.getsize(pdf)//1024}KB  '
          f'pptx {os.path.getsize(pptx)//1024}KB')
    return {'pdf': os.path.basename(pdf), 'pptx': os.path.basename(pptx), 'pages': n_pages}


def main():
    force = '--force' in sys.argv
    slugs = [a for a in sys.argv[1:] if not a.startswith('-')]
    if not slugs:
        slugs = [os.path.basename(os.path.dirname(p))
                 for p in glob.glob(os.path.join(ROOT, 'slides', '*', 'index.html'))
                 if 'hs-canvas' in open(p, encoding='utf-8').read()]
    if not os.path.exists(CHROME):
        raise SystemExit('headless-chrome not found at ' + CHROME)
    todo = sorted(slugs if force else [s for s in slugs if is_stale(s)])
    skipped = len(slugs) - len(todo)
    print(f'exporting {len(todo)} deck(s)' + (f', {skipped} already current' if skipped else ''))
    for s in todo:
        export(s)


if __name__ == '__main__':
    main()
